import io
import os
import tempfile
from datetime import datetime

import cv2
import pandas as pd
from PIL import Image

def hw_text(x):
    if x["hw_digits"]:
        return x["hw_digits"]
    return "none read" if x.get("hw_enabled") else "not read (handwriting reader switched off)"


DISCLAIMER = ("This report was produced by an AI-assisted image reading tool. It is intended to support, not replace, "
              "reading by trained laboratory staff. Every result must be confirmed by a person before it is used. "
              "It is not a clinical validation or an authorization for autonomous diagnostic use.")


def summary_counts(results):
    n = len(results)
    c = {"photos": n,
         "positive": sum(r["row"]["photo_result"] == "positive" for r in results),
         "negative": sum(r["row"]["photo_result"] == "negative" for r in results),
         "invalid": sum(r["row"]["photo_result"] == "invalid" for r in results),
         "review": sum(r["row"]["needs_review"] for r in results)}
    return c


def results_frame(results):
    rows = []
    for r in results:
        x = r["row"]
        rows.append({
            "File": r["name"],
            "Result": x["photo_result"].upper(),
            "Devices": x["n_devices"],
            "Device results": ", ".join(x["device_results"]),
            "Specimen ID (barcode)": x["specimen_id"] or "not read",
            "Handwritten ID": hw_text(x),
            "Test type": x["test_type"],
            "Needs review": "Yes" if x["needs_review"] else "No",
            "Notes": " ".join(x["reasons"]),
        })
    return pd.DataFrame(rows)


def annotated_jpeg(img_rgb, max_w=1400, quality=85):
    im = Image.fromarray(img_rgb)
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=quality)
    return buf.getvalue()


def build_docx(results, title="HBsAg Rapid Test Photograph Report", operator=""):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)

    def heading(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        r.font.name = "Times New Roman"
        return p

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(datetime.now().strftime("%d %B %Y, %H:%M")).italic = True
    if operator:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"Prepared by: {operator}")

    c = summary_counts(results)
    heading("1. Summary")
    doc.add_paragraph(f"{c['photos']} photograph(s) were analysed: {c['positive']} positive, {c['negative']} negative and {c['invalid']} invalid. "
                      f"{c['review']} photograph(s) need a person to check them before the result is used.")
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["File", "Result", "Specimen ID", "Test type", "Review"]):
        t.rows[0].cells[i].text = h
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for r in results:
        x = r["row"]
        cells = t.add_row().cells
        for i, v in enumerate([r["name"], x["photo_result"].upper(), x["specimen_id"] or "not read", x["test_type"], "Yes" if x["needs_review"] else "No"]):
            cells[i].text = v
    doc.add_paragraph()

    heading("2. Photograph Details")
    for k, r in enumerate(results, 1):
        x = r["row"]
        p = doc.add_paragraph()
        run = p.add_run(f"2.{k} {r['name']}: {x['photo_result'].upper()}")
        run.bold = True
        p.paragraph_format.keep_with_next = True
        im = Image.open(io.BytesIO(r["jpeg"]))
        k = min(5.6 / im.width, 4.6 / im.height)
        doc.add_picture(io.BytesIO(r["jpeg"]), width=Inches(im.width * k), height=Inches(im.height * k))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        facts = [("Devices in the photograph", str(x["n_devices"])),
                 ("Device results", ", ".join(d.upper() for d in x["device_results"]) or "none found"),
                 ("Specimen ID from barcode", (x["specimen_id"] + (f" ({x['specimen_id_source']})" if x["specimen_id_source"] else "")) if x["specimen_id"] else "not read"),
                 ("Handwritten ID", hw_text(x)),
                 ("Test type", x["test_type"]),
                 ("Needs review", "Yes" if x["needs_review"] else "No")]
        for a, b in facts:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = a, b
            cells[0].paragraphs[0].runs[0].bold = True
        if x["reasons"]:
            p = doc.add_paragraph()
            p.add_run("Points to check:").bold = True
            for reason in x["reasons"]:
                doc.add_paragraph(reason, style="List Bullet")
        doc.add_paragraph()

    heading("3. Important Notice")
    doc.add_paragraph(DISCLAIMER)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(results, title="HBsAg Rapid Test Photograph Report", operator=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image as RLImage, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Times-Roman", fontSize=11, leading=14)
    small = ParagraphStyle("s", parent=body, fontSize=9.5, leading=12)
    h1 = ParagraphStyle("h1", parent=body, fontName="Times-Bold", fontSize=14, textColor=colors.HexColor("#2E75B6"), spaceBefore=10, spaceAfter=6)
    ttl = ParagraphStyle("t", parent=body, fontName="Times-Bold", fontSize=18, alignment=1, textColor=colors.HexColor("#1F4E79"), spaceAfter=4)
    cen = ParagraphStyle("c", parent=body, alignment=1, textColor=colors.HexColor("#555555"))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    story = [Paragraph(title, ttl), Paragraph(datetime.now().strftime("%d %B %Y, %H:%M"), cen)]
    if operator:
        story.append(Paragraph(f"Prepared by: {operator}", cen))
    c = summary_counts(results)
    story += [Paragraph("1. Summary", h1),
              Paragraph(f"{c['photos']} photograph(s) were analysed: {c['positive']} positive, {c['negative']} negative and {c['invalid']} invalid. "
                        f"{c['review']} photograph(s) need a person to check them before the result is used.", body), Spacer(1, 8)]
    data = [["File", "Result", "Specimen ID", "Test type", "Review"]]
    for r in results:
        x = r["row"]
        data.append([Paragraph(r["name"], small), x["photo_result"].upper(), x["specimen_id"] or "not read", x["test_type"], "Yes" if x["needs_review"] else "No"])
    t = Table(data, colWidths=[2.2 * inch, 0.9 * inch, 1.4 * inch, 1.1 * inch, 0.8 * inch], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")), ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                           ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"), ("FONTSIZE", (0, 0), (-1, -1), 9.5), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story += [t, Paragraph("2. Photograph Details", h1)]
    for k, r in enumerate(results, 1):
        x = r["row"]
        im = Image.open(io.BytesIO(r["jpeg"]))
        k = min(6.2 * inch / im.width, 4.6 * inch / im.height)
        img = RLImage(io.BytesIO(r["jpeg"]), width=im.width * k, height=im.height * k)
        facts = [["Devices in the photograph", str(x["n_devices"])],
                 ["Device results", ", ".join(d.upper() for d in x["device_results"]) or "none found"],
                 ["Specimen ID from barcode", (x["specimen_id"] + (f" ({x['specimen_id_source']})" if x["specimen_id_source"] else "")) if x["specimen_id"] else "not read"],
                 ["Handwritten ID", hw_text(x)], ["Test type", x["test_type"]], ["Needs review", "Yes" if x["needs_review"] else "No"]]
        ft = Table(facts, colWidths=[2.4 * inch, 3.8 * inch])
        ft.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")), ("FONTNAME", (0, 0), (0, -1), "Times-Bold"),
                                ("FONTNAME", (1, 0), (1, -1), "Times-Roman"), ("FONTSIZE", (0, 0), (-1, -1), 9.5)]))
        block = [Paragraph(f"<b>2.{k} {r['name']}: {x['photo_result'].upper()}</b>", body), Spacer(1, 4), img, Spacer(1, 4), ft]
        if x["reasons"]:
            block.append(Spacer(1, 4))
            block.append(Paragraph("<b>Points to check:</b> " + " ".join(x["reasons"]), small))
        block.append(Spacer(1, 12))
        story.append(KeepTogether(block))
    story += [Paragraph("3. Important Notice", h1), Paragraph(DISCLAIMER, body)]
    doc.build(story)
    return buf.getvalue()
